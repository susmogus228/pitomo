"""Bounded admin analytics exports with explicit, non-secret field projections."""
from __future__ import annotations

import csv
import json
import logging
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
from io import BytesIO, StringIO
from threading import BoundedSemaphore
from typing import Annotated, Literal
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi import APIRouter, HTTPException, Query, Response
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import exists, func, or_, select
from sqlalchemy.orm import Session

from app.auth.dependencies import AdminUser
from app.catalog.pricing import active_discount_percent
from app.catalog.search import product_search_criteria
from .order_filters import (
    DeliveryStageFilter, DeliveryIncidentFilter, OrderQueue, order_criteria,
)
from app.db.dependencies import SessionDependency
from app.db.models import (
    OrderItemORM, OrderORM, ProductORM, ProductVariantORM, ReviewORM, ReviewPhotoORM,
    ReviewReplyORM, StockReplenishmentORM, UserORM,
)

router = APIRouter(prefix='/api/v1/admin/exports', tags=['Экспорт аналитики'])
logger = logging.getLogger('petshop.audit')
Dataset = Literal['overview', 'orders', 'products', 'replenishments', 'users', 'reviews',
                  'discounts', 'all']
ExportFormat = Literal['xlsx', 'csv']
KAZAKHSTAN_TZ = timezone(timedelta(hours=5))
MAX_EXPORT_ROWS = 100_000
MAX_DATA_BYTES = 20 * 1024 * 1024
MAX_EXPORT_BYTES = 50 * 1024 * 1024
_EXPORT_SLOTS = BoundedSemaphore(2)
_CONTROL_CHARACTERS = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
_PHOTO_FILENAME = re.compile(r'[a-f0-9]{32}\.webp')


def _too_large() -> HTTPException:
    return HTTPException(413, 'Выгрузка слишком большая. Сократите период или уточните фильтры.')


@dataclass
class ExportBudget:
    rows: int = 0
    data_bytes: int = 0

    def account(self, row: tuple) -> None:
        self.rows += 1
        self.data_bytes += sum(len(value.encode('utf-8', errors='replace'))
                               if isinstance(value, str) else 32 for value in row)
        if self.rows > MAX_EXPORT_ROWS or self.data_bytes > MAX_DATA_BYTES:
            raise _too_large()


@dataclass
class ExportTable:
    name: str
    columns: tuple[str, ...]
    rows: list[tuple] = field(default_factory=list)

    def append(self, values, budget: ExportBudget) -> None:
        row = tuple(values)
        budget.account(row)
        self.rows.append(row)


@dataclass(frozen=True)
class ExportFilters:
    date_from: date | None
    date_to: date | None
    q: str | None
    status: str | None
    payment_status: str | None
    delivery_method: str | None
    user_id: int | None
    variant_id: int | None
    unanswered_only: bool
    include_contacts: bool
    delivery_stage: str | None = None
    delivery_incident: str | None = None
    received: bool | None = None
    stale_hours: int | None = None
    queue: str | None = None
    stock_max: int | None = None
    visibility: str = 'all'

    def period(self, column) -> list:
        criteria = []
        if self.date_from is not None and self.date_from != date.min:
            criteria.append(column >= datetime.combine(self.date_from, time.min,
                                                       KAZAKHSTAN_TZ).astimezone(timezone.utc))
        if self.date_to is not None:
            end = datetime.combine(self.date_to, time.min, KAZAKHSTAN_TZ)
            # date.max cannot be incremented; every representable date is included instead.
            if self.date_to != date.max:
                criteria.append(column < (end + timedelta(days=1)).astimezone(timezone.utc))
        return criteria

    def orders(self) -> list:
        return self.period(OrderORM.created_at) + order_criteria(
            status=self.status, payment_status=self.payment_status,
            delivery_method=self.delivery_method, user_id=self.user_id,
            delivery_stage=self.delivery_stage, delivery_incident=self.delivery_incident,
            received=self.received, stale_hours=self.stale_hours, queue=self.queue)

    def products(self) -> list:
        criteria = product_search_criteria(self.q)
        if self.stock_max is not None:
            criteria.append(select(ProductVariantORM.id).where(
                ProductVariantORM.product_id == ProductORM.id,
                ProductVariantORM.stock <= self.stock_max).correlate(ProductORM).exists())
        return criteria


def _read_table(session: Session, statement, name: str, columns: tuple[str, ...],
                budget: ExportBudget) -> ExportTable:
    table = ExportTable(name, columns)
    rows = session.execute(statement.limit(MAX_EXPORT_ROWS + 1)
                           .execution_options(yield_per=512)).mappings()
    try:
        for row in rows:
            table.append((row[column] for column in columns), budget)
    finally:
        rows.close()
    return table


def _order_tables(session: Session, filters: ExportFilters, budget: ExportBudget,
                  *, include_history: bool = False):
    order_columns = ['id', 'user_id', 'status', 'created_at', 'delivery_method', 'payment_method',
                     'payment_status', 'paid_at', 'received_at', 'delivery_stage',
                     'delivery_incident', 'delivery_version', 'delivery_updated_at', 'currency', 'total']
    projected_order = [getattr(OrderORM, column).label('order_' + column)
                       for column in order_columns[:-2]]
    if include_history:
        projected_order.append(OrderORM.delivery_history.label('order_delivery_history'))
    if filters.include_contacts:
        contact_columns = ['contact_name', 'phone', 'address', 'comment', 'delivery_lat', 'delivery_lon',
                           'delivery_incident_note']
        order_columns += ['customer_email', 'customer_first_name', 'customer_last_name', *contact_columns]
        projected_order += [UserORM.email.label('order_customer_email'),
                            UserORM.first_name.label('order_customer_first_name'),
                            UserORM.last_name.label('order_customer_last_name')]
        projected_order += [getattr(OrderORM, column).label('order_' + column)
                            for column in contact_columns]
    item_columns = ('id', 'order_id', 'user_id', 'product_id', 'product_name', 'variant_id', 'sku',
                    'variant_label', 'quantity', 'currency', 'unit_price', 'original_unit_price',
                    'line_total', 'line_discount')
    projected_item = [getattr(OrderItemORM, column).label('item_' + column)
                      for column in item_columns if column not in {'user_id', 'line_total', 'line_discount'}]
    statement = select(*projected_order, *projected_item).select_from(OrderORM).outerjoin(
        OrderItemORM, OrderItemORM.order_id == OrderORM.id).where(*filters.orders()).order_by(
        OrderORM.id, OrderItemORM.id)
    if filters.include_contacts:
        statement = statement.outerjoin(UserORM, UserORM.id == OrderORM.user_id)
    # A single statement keeps an order's status and its historical lines in the same snapshot.
    rows = session.execute(statement.limit(MAX_EXPORT_ROWS + 1)
                           .execution_options(yield_per=512)).mappings()
    orders: dict[int, dict] = {}
    items = ExportTable('OrderItems', item_columns)
    try:
        for row in rows:
            order_id = row['order_id']
            if order_id not in orders:
                order = {column: row['order_' + column]
                         for column in order_columns if column not in {'currency', 'total'}}
                order['totals'] = defaultdict(Decimal)
                if include_history:
                    order['history'] = row['order_delivery_history']
                orders[order_id] = order
                # Account the order now; calculated amounts do not add unbounded strings later.
                budget.account(tuple(order[column] for column in order_columns
                                     if column not in {'currency', 'total'}))
            if row['item_id'] is not None:
                item = {column: row['item_' + column] for column in item_columns
                        if column not in {'user_id', 'line_total', 'line_discount'}}
                item['user_id'] = row['order_user_id']
                item['line_total'] = item['unit_price'] * item['quantity']
                original = item['original_unit_price']
                item['line_discount'] = ((original - item['unit_price']) * item['quantity']
                                         if original is not None else Decimal('0.00'))
                orders[order_id]['totals'][item['currency']] += item['line_total']
                items.append((item[column] for column in item_columns), budget)
    finally:
        rows.close()
    order_table = ExportTable('Orders', tuple(order_columns))
    for order in orders.values():
        amounts = order['totals']
        order['currency'] = next(iter(amounts)) if len(amounts) == 1 else ('MULTI' if amounts else 'KZT')
        order['total'] = next(iter(amounts.values())) if len(amounts) == 1 else (None if amounts else Decimal('0.00'))
        order_table.rows.append(tuple(order[column] for column in order_columns))
    return order_table, items, orders


def _delivery_events(orders: dict, filters: ExportFilters, budget: ExportBudget) -> ExportTable:
    columns = ('order_id', 'event_index', 'event_type', 'stage', 'incident', 'occurred_at')
    if filters.include_contacts:
        columns += ('note',)
    table = ExportTable('DeliveryEvents', columns)
    # Reuse the same statement snapshot as Orders/OrderItems, including concurrent state changes.
    for order_id, order in orders.items():
        for index, event in enumerate(order.get('history') or [], 1):
            # Explicit projections keep addresses, arbitrary JSON fields and private notes out.
            occurred_at = event.get('occurred_at')
            if isinstance(occurred_at, str):
                occurred_at = datetime.fromisoformat(occurred_at)
            values = (order_id, index, event.get('event_type'), event.get('stage'),
                      event.get('incident'), occurred_at)
            if filters.include_contacts:
                values += (event.get('note'),)
            table.append(values, budget)
    return table


def _catalog_tables(session: Session, filters: ExportFilters, budget: ExportBudget):
    columns = ('id', 'name', 'pet_type', 'brand', 'category', 'product_kind', 'image_url',
               'composition', 'material', 'discount_percent', 'discount_start_at', 'discount_end_at')
    products = _read_table(session, select(*(getattr(ProductORM, col) for col in columns)).where(
        *filters.products()).order_by(ProductORM.id), 'Products', columns, budget)
    variant_columns = ('id', 'product_id', 'sku', 'label', 'price', 'currency', 'stock')
    variants = _read_table(session, select(*(getattr(ProductVariantORM, col) for col in variant_columns))
        .join(ProductORM, ProductORM.id == ProductVariantORM.product_id).where(*filters.products())
        .order_by(ProductVariantORM.id), 'Variants', variant_columns, budget)
    return products, variants


def _discounts(session: Session, filters: ExportFilters, budget: ExportBudget,
               generated_at: datetime) -> ExportTable:
    columns = ('product_id', 'product_name', 'discount_percent', 'start_at', 'end_at')
    statement = select(ProductORM.id.label('product_id'), ProductORM.name.label('product_name'),
                       ProductORM.discount_percent, ProductORM.discount_start_at.label('start_at'),
                       ProductORM.discount_end_at.label('end_at')).where(
        ProductORM.discount_percent.is_not(None), *filters.products()).order_by(ProductORM.id)
    table = _read_table(session, statement, 'Discounts', columns, budget)
    table.columns += ('is_active_at_export',)
    table.rows = [(*row, active_discount_percent(row[2], row[3], row[4], generated_at) is not None)
                  for row in table.rows]
    return table


def _replenishments(session: Session, filters: ExportFilters, budget: ExportBudget) -> ExportTable:
    columns = ('id', 'variant_id', 'actor_id', 'product_name', 'variant_label', 'sku', 'quantity',
               'stock_before', 'stock_after', 'reason', 'created_at')
    criteria = filters.period(StockReplenishmentORM.created_at)
    if filters.variant_id is not None:
        criteria.append(StockReplenishmentORM.variant_id == filters.variant_id)
    if filters.q:
        lowercase = func.lower
        if session.get_bind().dialect.name == 'sqlite':
            # SQLite's built-in lower() handles ASCII only; test/local storage needs Cyrillic too.
            connection = session.connection().connection.driver_connection
            connection.create_function('petshop_export_lower', 1,
                                       lambda value: value.lower() if value else '', deterministic=True)
            lowercase = func.petshop_export_lower
        search = filters.q.strip().lower()
        criteria.append(or_(lowercase(StockReplenishmentORM.product_name).contains(search, autoescape=True),
                            lowercase(StockReplenishmentORM.sku).contains(search, autoescape=True)))
    # These immutable audit snapshots predate/depend on removed catalog variants as well.
    table = _read_table(session, select(*(getattr(StockReplenishmentORM, col) for col in columns))
        .where(*criteria).order_by(StockReplenishmentORM.id), 'Replenishments', columns, budget)
    return table


def _review_tables(session: Session, filters: ExportFilters, budget: ExportBudget):
    answered = exists(select(ReviewReplyORM.id).where(
        ReviewReplyORM.review_id == ReviewORM.id, ReviewReplyORM.is_pinned.is_(True),
        ReviewReplyORM.deleted_at.is_(None)).correlate(ReviewORM))
    criteria = [ReviewORM.deleted_at.is_(None), *filters.period(ReviewORM.created_at)]
    if filters.visibility != 'all':
        criteria.append(ReviewORM.hidden_at.is_(None) if filters.visibility == 'visible' else ReviewORM.hidden_at.is_not(None))
    if filters.unanswered_only:
        criteria.append(~answered)
    columns = ('id', 'product_id', 'order_id', 'user_id', 'rating', 'body', 'created_at', 'updated_at', 'hidden_at', 'edited_at',
               'has_official_reply', 'product_name')
    product_name = select(OrderItemORM.product_name).where(
        OrderItemORM.order_id == ReviewORM.order_id, OrderItemORM.product_id == ReviewORM.product_id
    ).correlate(ReviewORM).order_by(OrderItemORM.id).limit(1).scalar_subquery()
    reviews = _read_table(session, select(*(getattr(ReviewORM, col) for col in columns[:-2]),
        answered.label('has_official_reply'), product_name.label('product_name')).where(*criteria).order_by(ReviewORM.id),
        'Reviews', columns, budget)
    reply_columns = ('id', 'review_id', 'parent_reply_id', 'user_id', 'body', 'is_official',
                     'is_pinned', 'created_at', 'updated_at')
    replies = _read_table(session, select(*(getattr(ReviewReplyORM, col) for col in reply_columns))
        .join(ReviewORM, ReviewORM.id == ReviewReplyORM.review_id).where(
            *criteria, ReviewReplyORM.deleted_at.is_(None)).order_by(ReviewReplyORM.id),
        'ReviewReplies', reply_columns, budget)
    selected_review_ids = {row[0] for row in reviews.rows}
    # A newly posted review may appear in a later statement in READ COMMITTED mode.
    # Never export its replies without their selected parent review.
    replies.rows = [row for row in replies.rows if row[1] in selected_review_ids]
    selected_reply_ids = {row[0] for row in replies.rows}
    photos = session.execute(select(ReviewPhotoORM.review_id, ReviewPhotoORM.reply_id,
        ReviewPhotoORM.filename).outerjoin(ReviewReplyORM, ReviewReplyORM.id == ReviewPhotoORM.reply_id)
        .join(ReviewORM, or_(ReviewORM.id == ReviewPhotoORM.review_id,
                            ReviewORM.id == ReviewReplyORM.review_id)).where(
            *criteria, ReviewPhotoORM.deleted_at.is_(None),
            or_(ReviewPhotoORM.reply_id.is_(None), ReviewReplyORM.deleted_at.is_(None)))
        .order_by(ReviewPhotoORM.id).limit(MAX_EXPORT_ROWS + 1).execution_options(yield_per=512))
    photo_urls: dict[tuple[str, int], list[str]] = defaultdict(list)
    try:
        for photo in photos:
            if ((photo.review_id is not None and photo.review_id not in selected_review_ids)
                    or (photo.reply_id is not None and photo.reply_id not in selected_reply_ids)):
                continue
            if _PHOTO_FILENAME.fullmatch(photo.filename):
                key = ('review', photo.review_id) if photo.review_id is not None else ('reply', photo.reply_id)
                url = '/media/reviews/' + photo.filename
                budget.account((url,))
                photo_urls[key].append(url)
    finally:
        photos.close()
    for table, kind in ((reviews, 'review'), (replies, 'reply')):
        table.columns += ('photo_urls',)
        table.rows = [(*row, '\n'.join(photo_urls.get((kind, row[0]), ()))) for row in table.rows]
    return reviews, replies


def _category(order: dict) -> tuple[str, ...]:
    result = ['all', order['status']]
    if order['status'] == 'created' and order['received_at'] is not None:
        result.append('received')
    if order['status'] == 'created' and order['payment_method'] == 'demo' and order['payment_status'] == 'demo_paid':
        result.append('demo_paid')
    return tuple(result)


def _overview(orders: dict, items: ExportTable, budget: ExportBudget) -> ExportTable:
    columns = ('metric', 'product_id', 'product_name', 'currency', 'orders_count', 'quantity', 'amount')
    table = ExportTable('Overview', columns)
    totals = defaultdict(lambda: {'orders': set(), 'quantity': 0, 'amount': Decimal('0.00')})
    for item_row in items.rows:
        item = dict(zip(items.columns, item_row))
        order = orders[item['order_id']]
        for category in _category(order):
            total = totals[(category, None, None, item['currency'])]
            total['orders'].add(order['id'])
            total['quantity'] += item['quantity']
            total['amount'] += item['line_total']
        if order['status'] == 'created':
            total = totals[('created_product', item['product_id'], item['product_name'], item['currency'])]
            total['orders'].add(order['id'])
            total['quantity'] += item['quantity']
            total['amount'] += item['line_total']
    for category in ('all', 'created', 'cancelled', 'received', 'demo_paid'):
        count = sum(category in _category(order) for order in orders.values())
        table.append((category + '_orders_count', None, None, None, count, None, None), budget)
        if not any(key[0] == category for key in totals):
            table.append((category + '_order_value', None, None, 'KZT', 0, 0, Decimal('0.00')), budget)
    for (category, product_id, product_name, currency), total in totals.items():
        table.append((category + '_order_value', product_id, product_name, currency,
                      len(total['orders']), total['quantity'], total['amount']), budget)
    return table


def _users(session: Session, filters: ExportFilters, orders: dict,
           budget: ExportBudget) -> ExportTable:
    columns = ['id', 'role', 'is_active']
    if filters.include_contacts:
        columns += ['email', 'first_name', 'last_name']
    source = _read_table(session, select(*(getattr(UserORM, col) for col in columns))
        .order_by(UserORM.id), 'Users', tuple(columns), budget)
    totals = defaultdict(lambda: {'count': 0, 'created_count': 0, 'cancelled_count': 0,
        'created': Decimal('0.00'), 'cancelled': Decimal('0.00'), 'received': Decimal('0.00'),
        'demo_paid': Decimal('0.00')})
    currencies: dict[int, set[str]] = defaultdict(set)
    for order in orders.values():
        for currency, amount in (order['totals'] or {'KZT': Decimal('0.00')}).items():
            currencies[order['user_id']].add(currency)
            total = totals[(order['user_id'], currency)]
            total['count'] += 1
            total[order['status'] + '_count'] += 1
            for category in _category(order):
                if category != 'all':
                    total[category] += amount
    source_columns = tuple(columns)
    source.columns += ('currency', 'orders_count', 'created_orders_count', 'cancelled_orders_count',
                       'created_order_value', 'cancelled_order_value', 'received_order_value',
                       'demo_paid_order_value')
    result = []
    for row in source.rows:
        for index, currency in enumerate(sorted(currencies[row[0]] or {'KZT'})):
            total = totals[(row[0], currency)]
            values = (*row, currency, total['count'], total['created_count'], total['cancelled_count'],
                      total['created'], total['cancelled'], total['received'], total['demo_paid'])
            if index == 0:
                # The source user was already counted; budget only the additional statistics.
                budget.data_bytes += 32 * (len(source.columns) - len(source_columns))
            else:
                budget.account(values)
            result.append(values)
            if budget.data_bytes > MAX_DATA_BYTES:
                raise _too_large()
    source.rows = result
    return source


def _metadata(filters: ExportFilters, dataset: str, generated_at: datetime,
              tables: list[ExportTable], budget: ExportBudget) -> ExportTable:
    metadata = ExportTable('Metadata', ('key', 'value'))
    values = {
        'dataset': dataset, 'generated_at': generated_at.isoformat(), 'timezone': 'UTC+05:00',
        'date_from': filters.date_from.isoformat() if filters.date_from else '',
        'date_to_inclusive': filters.date_to.isoformat() if filters.date_to else '',
        'q': filters.q or '', 'status': filters.status or '',
        'payment_status': filters.payment_status or '', 'delivery_method': filters.delivery_method or '',
        'delivery_stage': filters.delivery_stage or '', 'delivery_incident': filters.delivery_incident or '',
        'received': filters.received if filters.received is not None else '',
        'stale_hours': filters.stale_hours if filters.stale_hours is not None else '',
        'queue': filters.queue or '',
        'stock_max': filters.stock_max if filters.stock_max is not None else '',
        'user_id': filters.user_id or '', 'variant_id': filters.variant_id or '',
        'unanswered_only': filters.unanswered_only, 'include_contacts': filters.include_contacts,
        'review_visibility': filters.visibility,
        'period_scope': 'Order/review/replenishment created_at; events/replies/photos belong to selected parents.',
        'order_filters_scope': 'Orders, OrderItems, DeliveryEvents, Overview and Users statistics.',
        'search_scope': 'Catalog name/SKU/brand and pet/category vocabulary; historical replenishment name or SKU.',
        'queue_semantics': 'active=created unreceived; incidents=active delivery regardless of payment; warehouse/in_transit/awaiting_receipt require paid demo or COD. received=created received.',
        'stale_semantics': 'Active delivery last updated (created_at fallback) at least stale_hours ago.',
        'delivery_events_semantics': 'History of selected orders; notes require include_contacts. Empty history sheet omitted.',
        'catalog_scope': 'Products/Variants/Discounts are current settings and stock at export time.',
        'stock_scope': 'stock_max selects products with any SKU at or below threshold; all sibling SKUs included.',
        'amount_semantics': 'Order values only. demo_paid is test payment; received COD can remain cod_pending.',
        'summary_overlap': 'Received and demo_paid are overlapping subsets of created. Do not add them.',
        'currency_semantics': 'Amounts are grouped by currency; multi-currency Orders total is empty.',
        'users_semantics': 'All users; zero for no selected orders. Composite key (id,currency).',
        'excluded_reviews': 'Deleted reviews/replies/photos are omitted.',
        'row_limit': MAX_EXPORT_ROWS,
    }
    values.update({'rows.' + table.name: len(table.rows) for table in tables})
    for item in values.items():
        metadata.append(item, budget)
    return metadata


def _collect(session: Session, filters: ExportFilters, dataset: str, generated_at: datetime):
    budget = ExportBudget()
    tables = []
    if dataset in {'orders', 'overview', 'users', 'all'}:
        orders_table, items_table, orders = _order_tables(
            session, filters, budget, include_history=dataset in {'orders', 'all'})
        if dataset in {'orders', 'all'}:
            tables += [orders_table, items_table]
            events = _delivery_events(orders, filters, budget)
            if events.rows:
                # Preserve the legacy ZIP shape when selected orders have no delivery history.
                tables.append(events)
        if dataset in {'overview', 'all'}:
            tables.append(_overview(orders, items_table, budget))
        if dataset in {'users', 'all'}:
            tables.append(_users(session, filters, orders, budget))
    if dataset in {'products', 'all'}:
        tables += _catalog_tables(session, filters, budget)
    if dataset in {'discounts', 'all'}:
        tables.append(_discounts(session, filters, budget, generated_at))
    if dataset in {'replenishments', 'all'}:
        tables.append(_replenishments(session, filters, budget))
    if dataset in {'reviews', 'all'}:
        tables += _review_tables(session, filters, budget)
    metadata = _metadata(filters, dataset, generated_at, tables, budget)
    return [metadata, *tables]


def _text(value: str) -> str:
    return _CONTROL_CHARACTERS.sub('', value).encode('utf-8', errors='replace').decode('utf-8')


def _kz_time(value: datetime) -> datetime:
    return (value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value).astimezone(KAZAKHSTAN_TZ)


def _xlsx(tables: list[ExportTable]) -> bytes:
    # Preflight prevents Excel's silent 32767-character cell truncation.
    for table in tables:
        for row in table.rows:
            if any(isinstance(value, str) and len(_text(value)) > 32767 for value in row):
                raise HTTPException(413, 'Поле превышает лимит Excel в 32767 символов. Выберите CSV.')
    workbook = Workbook(write_only=True)
    workbook.properties.title = 'Pitomo analytics'
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='3652F5')
    for table in tables:
        sheet = workbook.create_sheet(table.name)
        sheet.freeze_panes = 'A2'
        for index, column in enumerate(table.columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(40, max(16, len(column) + 2))
        header = []
        for column in table.columns:
            cell = WriteOnlyCell(sheet, value=column)
            cell.font, cell.fill = header_font, header_fill
            header.append(cell)
        sheet.append(header)
        for row in table.rows:
            cells = []
            for column, value in zip(table.columns, row):
                if isinstance(value, datetime):
                    value = _kz_time(value).replace(tzinfo=None)
                elif isinstance(value, str):
                    value = _text(value)
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = 's'  # Every user-supplied string remains text, never a formula.
                elif isinstance(value, Decimal):
                    cell.number_format = '0.0000000' if column in {'delivery_lat', 'delivery_lon'} else '#,##0.00'
                elif isinstance(value, datetime):
                    cell.number_format = 'yyyy-mm-dd hh:mm:ss'
                cells.append(cell)
            sheet.append(cells)
        sheet.auto_filter.ref = f'A1:{get_column_letter(len(table.columns))}{len(table.rows) + 1}'
    with BytesIO() as buffer:
        workbook.save(buffer)
        if buffer.tell() > MAX_EXPORT_BYTES:
            raise _too_large()
        return buffer.getvalue()


def _csv_value(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return _kz_time(value).isoformat()
    if isinstance(value, bool):
        return 'true' if value else 'false'
    if isinstance(value, str):
        value = _text(value)
        if value.startswith(('\t', '\r', '\n')) or value.lstrip().startswith(('=', '+', '-', '@')):
            value = "'" + value
    return value


def _csv(table: ExportTable) -> bytes:
    with StringIO(newline='') as buffer:
        writer = csv.writer(buffer)
        writer.writerow(table.columns)
        for row in table.rows:
            writer.writerow([_csv_value(value) for value in row])
            if buffer.tell() > MAX_EXPORT_BYTES:
                raise _too_large()
        value = buffer.getvalue().encode('utf-8-sig')
        if len(value) > MAX_EXPORT_BYTES:
            raise _too_large()
        return value


def _single_csv_table(dataset: str, tables: list[ExportTable]) -> ExportTable:
    selected = {table.name: table for table in tables}
    if dataset == 'products':
        products, variants = selected['Products'], selected['Variants']
        columns = products.columns + tuple('variant_' + column for column in variants.columns)
        result = ExportTable('Products', columns)
        by_product = defaultdict(list)
        for variant in variants.rows:
            by_product[variant[1]].append(variant)
        for product in products.rows:
            for variant in by_product[product[0]] or [tuple(None for _ in variants.columns)]:
                result.rows.append((*product, *variant))
        return result
    if dataset == 'reviews':
        reviews, replies = selected['Reviews'], selected['ReviewReplies']
        columns = ('record_type', 'id', 'review_id', 'parent_reply_id', 'product_id', 'product_name', 'order_id',
                   'user_id', 'rating', 'body', 'is_official', 'is_pinned', 'created_at', 'updated_at',
                   'hidden_at', 'edited_at', 'has_official_reply', 'photo_urls')
        result = ExportTable('Reviews', columns)
        parents = {row[0]: dict(zip(reviews.columns, row)) for row in reviews.rows}
        for table, kind in ((reviews, 'review'), (replies, 'reply')):
            for row in table.rows:
                entry = dict(zip(table.columns, row))
                entry['record_type'] = kind
                if kind == 'review':
                    entry['review_id'] = entry['id']
                else:
                    parent = parents[entry['review_id']]
                    entry['product_id'], entry['order_id'] = parent['product_id'], parent['order_id']
                    entry['product_name'] = parent['product_name']
                result.rows.append(tuple(entry.get(column) for column in columns))
        return result
    return tables[1]


def _csv_zip(tables: list[ExportTable]) -> bytes:
    with BytesIO() as buffer:
        total_uncompressed = 0
        with ZipFile(buffer, 'w', compression=ZIP_DEFLATED) as archive:
            for table in tables:
                data = _csv(table)
                total_uncompressed += len(data)
                if total_uncompressed > MAX_EXPORT_BYTES:
                    raise _too_large()
                archive.writestr(table.name + '.csv', data)
        if buffer.tell() > MAX_EXPORT_BYTES:
            raise _too_large()
        return buffer.getvalue()


@router.get('/{dataset}', summary='Скачать все отфильтрованные записи для аналитики',
    response_class=Response,
    responses={200: {'description': 'XLSX, CSV либо ZIP с отдельными CSV', 'content': {
        media: {'schema': {'type': 'string', 'format': 'binary'}} for media in (
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'text/csv', 'application/zip')}},
               401: {'description': 'Требуется вход'}, 403: {'description': 'Требуется администратор'},
               413: {'description': 'Сократите объём выгрузки'}, 422: {'description': 'Некорректные фильтры'},
               429: {'description': 'Другие выгрузки ещё готовятся'}})
def download_export(
    dataset: Dataset, admin_user: AdminUser, session: SessionDependency,
    format: ExportFormat = 'xlsx',
    date_from: date | None = None,
    date_to: date | None = None,
    q: Annotated[str | None, Query(max_length=100)] = None,
    status_filter: Annotated[Literal['created', 'cancelled'] | None, Query(alias='status')] = None,
    payment_status: Literal['pending', 'demo_paid', 'failed', 'cod_pending', 'cancelled'] | None = None,
    delivery_method: Literal['pickup', 'delivery'] | None = None,
    delivery_stage: DeliveryStageFilter | None = None,
    delivery_incident: DeliveryIncidentFilter | None = None,
    received: bool | None = None,
    stale_hours: Annotated[int | None, Query(ge=1, le=8760)] = None,
    queue: OrderQueue | None = None,
    stock_max: Annotated[int | None, Query(ge=0, le=2147483647)] = None,
    user_id: Annotated[int | None, Query(gt=0, le=2147483647)] = None,
    variant_id: Annotated[int | None, Query(gt=0, le=2147483647)] = None,
    unanswered_only: bool = False,
    visibility: Literal['all', 'visible', 'hidden'] = 'all',
    include_contacts: bool = False,
) -> Response:
    if date_from is not None and date_to is not None and date_from > date_to:
        raise HTTPException(422, 'Дата начала периода должна быть не позже даты окончания.')
    filters = ExportFilters(date_from, date_to, q, status_filter, payment_status, delivery_method,
                            user_id, variant_id, unanswered_only, include_contacts,
                            delivery_stage, delivery_incident, received, stale_hours, queue, stock_max, visibility)
    if not _EXPORT_SLOTS.acquire(blocking=False):
        raise HTTPException(429, 'Подождите завершения текущих выгрузок.', headers={'Retry-After': '3'})
    try:
        generated_at = datetime.now(timezone.utc)
        tables = _collect(session, filters, dataset, generated_at)
        if format == 'xlsx':
            payload, extension = _xlsx(tables), 'xlsx'
            media_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        elif dataset in {'orders', 'all'}:
            payload, extension = _csv_zip(tables), 'zip'
            media_type = 'application/zip'
        else:
            payload, extension = _csv(_single_csv_table(dataset, tables)), 'csv'
            media_type = 'text/csv; charset=utf-8'
        filename = f'pitomo-{dataset}-{generated_at:%Y%m%dT%H%M%SZ}.{extension}'
        logger.info(json.dumps({'event': 'admin_analytics_export', 'actor_id': admin_user.id,
            'dataset': dataset, 'format': format,
            'rows': {table.name: len(table.rows) for table in tables}}, ensure_ascii=True))
        return Response(payload, media_type=media_type, headers={
            'Content-Disposition': f'attachment; filename="{filename}"', 'Cache-Control': 'no-store',
            'X-Content-Type-Options': 'nosniff', 'X-Export-Timezone': 'UTC+05:00',
            'X-Export-Generated-At': generated_at.isoformat(),
        })
    finally:
        _EXPORT_SLOTS.release()
